from __future__ import annotations

import html
import io
import re
from typing import Any, Iterable

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from candidate_profile import CURRENT_ROLE, OLDER_EXPERIENCE, PROFILE

NAVY = "0C213D"
NAVY_2 = "142B4A"
GOLD = "C79A3B"
CHARCOAL = "252A31"
MID = "5E6773"
LIGHT = "F3F5F7"
WHITE = "FFFFFF"
LINE = "D9DEE5"


def _clean(value: Any) -> str:
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    return text


def _list(value: Any, limit: int | None = None) -> list[str]:
    if isinstance(value, list):
        items = value
    elif value:
        items = re.split(r"\n+|;", str(value))
    else:
        items = []
    output: list[str] = []
    seen: set[str] = set()
    for item in items:
        text = re.sub(r"^[\s•\-*\d.)]+", "", _clean(item)).strip()
        key = text.casefold()
        if text and key not in seen:
            seen.add(key)
            output.append(text)
        if limit and len(output) >= limit:
            break
    return output


def _shorten(text: str, max_words: int) -> str:
    words = _clean(text).split()
    if len(words) <= max_words:
        return " ".join(words)
    clipped = " ".join(words[:max_words]).rstrip(" ,;:-")
    return clipped + "."


def _headline(result: dict[str, Any], target_role: str = "") -> str:
    generated = _clean(result.get("target_headline"))
    if generated:
        return generated.upper()
    if target_role:
        role = re.sub(r"\s+", " ", target_role).strip().upper()
        if len(role) <= 62:
            return role
    return PROFILE["default_headline"]


def _profile_text(result: dict[str, Any]) -> str:
    value = _clean(result.get("executive_profile"))
    if not value:
        return (
            "Brand Experience and Partnerships leader with 10+ years in marketing and commercial leadership "
            "across Saudi Arabia and the GCC. Experienced in directing large-scale experiential programs, "
            "building strategic partner ecosystems, managing complex project portfolios, and translating "
            "brand activity into measurable pipeline and customer engagement outcomes."
        )
    return _shorten(value, 105)


def _current_bullets(result: dict[str, Any]) -> list[str]:
    generated = _list(result.get("current_experience"), limit=8)
    return [_shorten(item, 31) for item in generated]


def _capability_groups(result: dict[str, Any]) -> list[tuple[str, list[str]]]:
    groups: list[tuple[str, list[str]]] = []
    for item in _list(result.get("strategic_competencies"), limit=5):
        if ":" in item:
            heading, details = item.split(":", 1)
            skills = [x.strip() for x in details.split("|") if x.strip()]
        else:
            heading = "Core Capability"
            skills = [item]
        groups.append((_clean(heading).title(), skills[:5]))
    if groups:
        return groups
    return [
        ("Brand Experience", ["Experiential strategy", "Event delivery", "Brand storytelling"]),
        ("Partnerships", ["Strategic partnerships", "Vendor management", "Contract negotiation"]),
        ("Leadership", ["Cross-functional leadership", "Executive communication", "Team coaching"]),
        ("Commercial Delivery", ["Budget governance", "Project planning", "Performance measurement"]),
        ("Digital & CRM", ["Lifecycle marketing", "Campaign analytics", "Journey optimization"]),
    ]


def _career_highlights(result: dict[str, Any]) -> list[tuple[str, str]]:
    output: list[tuple[str, str]] = []
    for item in result.get("career_highlights", []) or []:
        if not isinstance(item, dict):
            continue
        value = _clean(item.get("value"))
        label = _clean(item.get("label"))
        if value and label:
            output.append((value, label))
        if len(output) == 5:
            break
    return output or list(PROFILE["metrics"])


def _signature_impact(result: dict[str, Any]) -> list[tuple[str, str]]:
    output: list[tuple[str, str]] = []
    for item in result.get("signature_impact", []) or []:
        if not isinstance(item, dict):
            continue
        label = _clean(item.get("label"))
        text = _clean(item.get("text"))
        if label and text:
            output.append((label.upper(), text))
        if len(output) == 4:
            break
    if output:
        return output
    return [
        ("FLAGSHIP EXPERIENCES", "BAIC and Maxus exhibitions in Jeddah."),
        ("COMMERCIAL PIPELINE", "3,000+ Sales Qualified Leads generated."),
        ("DELIVERY AT SCALE", "150+ employees, agencies and vendors led."),
        ("CUSTOMER ENGAGEMENT", "35% growth through CRM and lifecycle marketing."),
    ]


# -------------------------- DOCX helpers --------------------------

def _set_cell_shading(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_border(cell, **kwargs) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge, attrs in kwargs.items():
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        for key, value in attrs.items():
            node.set(qn(f"w:{key}"), str(value))


def _set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _remove_table_borders(table) -> None:
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def _set_table_fixed(table) -> None:
    tblPr = table._tbl.tblPr
    layout = tblPr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")


def _set_col_widths(table, widths: list[Inches]) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width
            tcPr = row.cells[idx]._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(int(width.inches * 1440)))
            tcW.set(qn("w:type"), "dxa")


def _font(run, size=None, bold=None, color=None, italic=None, name="Arial") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _para(p, before=0, after=0, line=1.0, keep=False, align=None) -> None:
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep
    if align is not None:
        p.alignment = align


def _clear_cell(cell) -> None:
    cell.text = ""
    _para(cell.paragraphs[0], after=0)



def _cell_paragraph(cell):
    if len(cell.paragraphs) == 1 and not cell.paragraphs[0].text:
        return cell.paragraphs[0]
    return cell.add_paragraph()


def _docx_text(cell, text: str, size=8.6, color=CHARCOAL, bold=False, after=3, line=1.05, italic=False):
    p = _cell_paragraph(cell)
    _para(p, after=after, line=line)
    r = p.add_run(text)
    _font(r, size=size, bold=bold, color=color, italic=italic)
    return p


def _docx_heading(cell, text: str, dark=False, top=7, bottom=5):
    p = _cell_paragraph(cell)
    _para(p, before=top, after=bottom, line=1, keep=True)
    r = p.add_run(text.upper())
    _font(r, 9.1, True, WHITE if dark else GOLD)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom_el = OxmlElement("w:bottom")
    bottom_el.set(qn("w:val"), "single")
    bottom_el.set(qn("w:sz"), "8")
    bottom_el.set(qn("w:space"), "3")
    bottom_el.set(qn("w:color"), GOLD if dark else LINE)
    pBdr.append(bottom_el)
    pPr.append(pBdr)
    return p


def _docx_bullet(cell, text: str, dark=False, size=8.3, after=2.5):
    p = _cell_paragraph(cell)
    _para(p, after=after, line=1.04)
    p.paragraph_format.left_indent = Inches(0.13)
    p.paragraph_format.first_line_indent = Inches(-0.13)
    r = p.add_run("•")
    _font(r, size + 0.2, True, GOLD)
    r2 = p.add_run("  " + text)
    _font(r2, size, False, WHITE if dark else CHARCOAL)
    return p


def _docx_label(cell, text: str, dark=True):
    p = _cell_paragraph(cell)
    _para(p, before=2, after=1, line=1, keep=True)
    r = p.add_run(text.upper())
    _font(r, 7.1, True, GOLD if dark else MID)
    return p


def _docx_header(doc: Document, headline: str, compact=False) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _remove_table_borders(table)
    _set_table_fixed(table)
    cell = table.cell(0, 0)
    _set_cell_shading(cell, NAVY)
    _set_cell_margins(cell, top=155 if not compact else 85, start=250, bottom=145 if not compact else 80, end=250)
    _clear_cell(cell)
    p = cell.paragraphs[0]
    _para(p, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(PROFILE["name"])
    _font(r, 22.5 if not compact else 15.5, True, WHITE)
    p2 = cell.add_paragraph()
    _para(p2, after=4 if not compact else 2, align=WD_ALIGN_PARAGRAPH.CENTER)
    r2 = p2.add_run(headline)
    _font(r2, 8.1 if not compact else 7.1, True, GOLD)
    p3 = cell.add_paragraph()
    _para(p3, align=WD_ALIGN_PARAGRAPH.CENTER)
    r3 = p3.add_run(f'{PROFILE["location"]}   |   {PROFILE["phone"]}   |   {PROFILE["email"]}')
    _font(r3, 7.8 if not compact else 7.0, False, "DDE5EE")


def _docx_metric(cell, value: str, label: str) -> None:
    _set_cell_shading(cell, LIGHT)
    _set_cell_border(
        cell,
        top={"val": "single", "sz": "22", "color": GOLD},
        bottom={"val": "single", "sz": "5", "color": LINE},
        left={"val": "single", "sz": "5", "color": LINE},
        right={"val": "single", "sz": "5", "color": LINE},
    )
    _set_cell_margins(cell, top=95, start=75, bottom=85, end=75)
    _clear_cell(cell)
    p = cell.paragraphs[0]
    _para(p, after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(value)
    _font(r, 15.0, True, NAVY)
    p2 = cell.add_paragraph()
    _para(p2, align=WD_ALIGN_PARAGRAPH.CENTER)
    r2 = p2.add_run(label.upper())
    _font(r2, 6.4, True, MID)


def _docx_experience(cell, title: str, company: str, location: str, dates: str, bullets: Iterable[str], highlight=False) -> None:
    p = _cell_paragraph(cell)
    _para(p, before=4 if highlight else 6, after=1, keep=True)
    r = p.add_run(title.upper())
    _font(r, 10.1 if highlight else 9.6, True, NAVY)
    p2 = cell.add_paragraph()
    _para(p2, after=2.5, keep=True)
    r2 = p2.add_run(company)
    _font(r2, 8.5, True, GOLD)
    r3 = p2.add_run(f"  |  {location}  |  {dates}")
    _font(r3, 7.9, False, MID)
    for bullet in bullets:
        _docx_bullet(cell, bullet, dark=False, size=8.15, after=2.2)


def generate_premium_docx(result: dict[str, Any], target_role: str = "") -> bytes:
    headline = _headline(result, target_role)
    profile_text = _profile_text(result)
    current_bullets = _current_bullets(result)
    capability_groups = _capability_groups(result)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.34)
    sec.bottom_margin = Inches(0.34)
    sec.left_margin = Inches(0.38)
    sec.right_margin = Inches(0.38)
    sec.header_distance = Inches(0.15)
    sec.footer_distance = Inches(0.15)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    doc.styles["Normal"].font.size = Pt(9)
    doc.styles["Normal"].paragraph_format.space_after = Pt(0)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _para(footer)
    rr = footer.add_run("ABDELRHMAN EL SHISHINY  •  EXECUTIVE RESUME")
    _font(rr, 7.0, False, MID)

    # Page 1
    _docx_header(doc, headline, compact=False)
    metrics = doc.add_table(rows=1, cols=5)
    metrics.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_fixed(metrics)
    _set_col_widths(metrics, [Inches(1.48)] * 5)
    _remove_table_borders(metrics)
    for idx, (value, label) in enumerate(_career_highlights(result)):
        _docx_metric(metrics.cell(0, idx), value, label)

    spacer = doc.add_paragraph()
    _para(spacer, before=1, after=1)

    outer = doc.add_table(rows=1, cols=2)
    outer.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_fixed(outer)
    _set_col_widths(outer, [Inches(2.05), Inches(5.32)])
    _remove_table_borders(outer)
    left, main = outer.rows[0].cells
    _set_cell_shading(left, NAVY_2)
    _set_cell_margins(left, top=140, start=150, bottom=120, end=145)
    _set_cell_margins(main, top=75, start=205, bottom=75, end=75)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    main.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    _clear_cell(left)
    _clear_cell(main)

    _docx_heading(left, "Executive Value", dark=True, top=0, bottom=6)
    for item in PROFILE["executive_value"]:
        _docx_bullet(left, item, dark=True, size=7.85, after=3.5)

    _docx_heading(left, "Selected Partnerships", dark=True, top=7, bottom=5)
    for label, value in PROFILE["selected_partnerships"]:
        _docx_label(left, label, dark=True)
        _docx_text(left, value, size=7.9, color=WHITE, after=3.5, line=1.04)

    _docx_heading(left, "Education", dark=True, top=7, bottom=5)
    for label, value in PROFILE["education"]:
        _docx_label(left, label, dark=True)
        _docx_text(left, value, size=7.9, color=WHITE, after=4, line=1.05)

    _docx_heading(left, "Languages", dark=True, top=7, bottom=5)
    for item in PROFILE["languages"]:
        _docx_text(left, item, size=7.9, color=WHITE, after=1.8)

    _docx_heading(main, "Executive Profile", top=0, bottom=5)
    _docx_text(main, profile_text, size=8.8, after=6, line=1.10)

    _docx_heading(main, "Signature Impact", top=2, bottom=5)
    impact = main.add_table(rows=1, cols=1)
    impact.alignment = WD_TABLE_ALIGNMENT.LEFT
    _set_table_fixed(impact)
    _remove_table_borders(impact)
    box = impact.cell(0, 0)
    _set_cell_shading(box, LIGHT)
    _set_cell_border(
        box,
        top={"val": "single", "sz": "18", "color": GOLD},
        bottom={"val": "single", "sz": "4", "color": LINE},
        left={"val": "single", "sz": "4", "color": LINE},
        right={"val": "single", "sz": "4", "color": LINE},
    )
    _set_cell_margins(box, top=70, start=105, bottom=65, end=105)
    _clear_cell(box)
    for label, text in _signature_impact(result):
        p = box.add_paragraph() if box.paragraphs and box.paragraphs[0].text else box.paragraphs[0]
        _para(p, after=1.7)
        r = p.add_run(label + "  ")
        _font(r, 6.8, True, GOLD)
        r2 = p.add_run(text)
        _font(r2, 7.55, False, CHARCOAL)

    _docx_heading(main, "Professional Experience", top=6, bottom=4)
    _docx_experience(
        main,
        CURRENT_ROLE["title"],
        CURRENT_ROLE["company"],
        CURRENT_ROLE["location"],
        CURRENT_ROLE["dates"],
        current_bullets,
        highlight=True,
    )

    page_break = doc.add_paragraph()
    page_break.add_run().add_break(WD_BREAK.PAGE)

    # Page 2
    _docx_header(doc, headline, compact=True)
    outer2 = doc.add_table(rows=1, cols=2)
    outer2.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_fixed(outer2)
    _set_col_widths(outer2, [Inches(2.05), Inches(5.32)])
    _remove_table_borders(outer2)
    left2, main2 = outer2.rows[0].cells
    _set_cell_shading(left2, NAVY_2)
    _set_cell_margins(left2, top=135, start=150, bottom=115, end=145)
    _set_cell_margins(main2, top=75, start=205, bottom=75, end=75)
    left2.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    main2.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    _clear_cell(left2)
    _clear_cell(main2)

    _docx_heading(left2, "Core Capabilities", dark=True, top=0, bottom=5)
    for label, skills in capability_groups:
        _docx_label(left2, label, dark=True)
        _docx_text(left2, "\n".join(skills), size=7.55, color=WHITE, after=4.5, line=1.03)

    _docx_heading(left2, "Tools", dark=True, top=6, bottom=5)
    _docx_text(left2, "\n".join(PROFILE["tools"]), size=7.75, color=WHITE, after=4, line=1.05)

    selected_courses = _list(result.get("selected_completed_courses"), limit=6)
    development = selected_courses or PROFILE["professional_development"]
    _docx_heading(left2, "Professional Development", dark=True, top=6, bottom=5)
    for item in development:
        _docx_bullet(left2, item, dark=True, size=7.45, after=2.1)

    _docx_heading(main2, "Earlier Professional Experience", top=0, bottom=4)
    for exp in OLDER_EXPERIENCE:
        _docx_experience(
            main2,
            exp["title"],
            exp["company"],
            exp["location"],
            exp["dates"],
            exp["bullets"],
        )

    _docx_heading(main2, "Leadership Positioning", top=7, bottom=4)
    position = _clean(result.get("current_role_positioning_line")) or (
        "Best aligned to Brand Experience, Strategic Partnerships, Commercial Partnerships, "
        "Marketing Programs and Experiential Marketing leadership roles."
    )
    _docx_text(main2, _shorten(position, 42), size=8.4, after=0, line=1.08, italic=True)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# -------------------------- PDF helpers --------------------------

def _esc(value: str) -> str:
    return html.escape(str(value or ""), quote=False)


def _pdf_styles() -> dict[str, ParagraphStyle]:
    return {
        "header_name": ParagraphStyle("header_name", fontName="Helvetica-Bold", fontSize=19, leading=21, textColor=colors.white, alignment=TA_CENTER, spaceAfter=2),
        "header_title": ParagraphStyle("header_title", fontName="Helvetica-Bold", fontSize=7.6, leading=9, textColor=colors.HexColor("#C79A3B"), alignment=TA_CENTER, spaceAfter=3),
        "header_contact": ParagraphStyle("header_contact", fontName="Helvetica", fontSize=7.3, leading=9, textColor=colors.HexColor("#DDE5EE"), alignment=TA_CENTER),
        "section": ParagraphStyle("section", fontName="Helvetica-Bold", fontSize=9.2, leading=11, textColor=colors.HexColor("#C79A3B"), spaceBefore=2, spaceAfter=4),
        "section_dark": ParagraphStyle("section_dark", fontName="Helvetica-Bold", fontSize=8.7, leading=10, textColor=colors.white, spaceBefore=2, spaceAfter=4),
        "body": ParagraphStyle("body", fontName="Helvetica", fontSize=8.15, leading=10.1, textColor=colors.HexColor("#252A31"), spaceAfter=3),
        "body_dark": ParagraphStyle("body_dark", fontName="Helvetica", fontSize=7.45, leading=9.3, textColor=colors.white, spaceAfter=2.5),
        "label_dark": ParagraphStyle("label_dark", fontName="Helvetica-Bold", fontSize=6.7, leading=8, textColor=colors.HexColor("#C79A3B"), spaceBefore=2, spaceAfter=1),
        "bullet": ParagraphStyle("bullet", fontName="Helvetica", fontSize=7.8, leading=9.7, textColor=colors.HexColor("#252A31"), leftIndent=8, firstLineIndent=-6, spaceAfter=1.9),
        "bullet_dark": ParagraphStyle("bullet_dark", fontName="Helvetica", fontSize=7.3, leading=9.1, textColor=colors.white, leftIndent=7, firstLineIndent=-5, spaceAfter=2.2),
        "role": ParagraphStyle("role", fontName="Helvetica-Bold", fontSize=9.1, leading=10.5, textColor=colors.HexColor("#0C213D"), spaceBefore=3, spaceAfter=1),
        "company": ParagraphStyle("company", fontName="Helvetica-Bold", fontSize=7.8, leading=9, textColor=colors.HexColor("#C79A3B"), spaceAfter=2),
        "meta": ParagraphStyle("meta", fontName="Helvetica", fontSize=7.1, leading=8.5, textColor=colors.HexColor("#5E6773")),
        "impact": ParagraphStyle("impact", fontName="Helvetica", fontSize=7.25, leading=8.9, textColor=colors.HexColor("#252A31"), spaceAfter=1.4),
        "metric_value": ParagraphStyle("metric_value", fontName="Helvetica-Bold", fontSize=12, leading=13, textColor=colors.HexColor("#0C213D"), alignment=TA_CENTER, spaceAfter=1),
        "metric_label": ParagraphStyle("metric_label", fontName="Helvetica-Bold", fontSize=5.7, leading=7, textColor=colors.HexColor("#5E6773"), alignment=TA_CENTER),
        "position": ParagraphStyle("position", fontName="Helvetica-Oblique", fontSize=7.9, leading=9.7, textColor=colors.HexColor("#252A31")),
    }


def _pdf_header(styles, headline: str, compact=False) -> Table:
    name_style = styles["header_name"]
    title_style = styles["header_title"]
    contact_style = styles["header_contact"]
    if compact:
        name_style = ParagraphStyle("header_name_compact", parent=name_style, fontSize=13.5, leading=15)
        title_style = ParagraphStyle("header_title_compact", parent=title_style, fontSize=6.6, leading=8, spaceAfter=1)
        contact_style = ParagraphStyle("header_contact_compact", parent=contact_style, fontSize=6.4, leading=7.5)
    data = [[
        Paragraph(_esc(PROFILE["name"]), name_style),
    ], [
        Paragraph(_esc(headline), title_style),
    ], [
        Paragraph(_esc(f'{PROFILE["location"]} | {PROFILE["phone"]} | {PROFILE["email"]}'), contact_style),
    ]]
    table = Table(data, colWidths=[194 * mm], hAlign="CENTER")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#0C213D")),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
        ("RIGHTPADDING", (0, 0), (-1, -1), 10),
        ("TOPPADDING", (0, 0), (-1, 0), 7 if not compact else 3),
        ("BOTTOMPADDING", (0, -1), (-1, -1), 7 if not compact else 3),
        ("TOPPADDING", (0, 1), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, 1), 1),
    ]))
    return table


def _pdf_metric_card(value: str, label: str, styles) -> Table:
    table = Table([
        [Paragraph(_esc(value), styles["metric_value"])],
        [Paragraph(_esc(label.upper()), styles["metric_label"])],
    ], colWidths=[38.4 * mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F3F5F7")),
        ("BOX", (0, 0), (-1, -1), 0.45, colors.HexColor("#D9DEE5")),
        ("LINEABOVE", (0, 0), (-1, 0), 1.6, colors.HexColor("#C79A3B")),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
        ("TOPPADDING", (0, 1), (-1, 1), 0),
        ("BOTTOMPADDING", (0, 1), (-1, 1), 4),
    ]))
    return table


def _pdf_section(title: str, styles, dark=False) -> list:
    style = styles["section_dark" if dark else "section"]
    line_color = colors.HexColor("#C79A3B" if dark else "#D9DEE5")
    return [
        Paragraph(_esc(title.upper()), style),
        Table([[""]], colWidths=[47 * mm if dark else 134 * mm], rowHeights=[0.7], style=TableStyle([("BACKGROUND", (0, 0), (-1, -1), line_color)])),
        Spacer(1, 2),
    ]


def _pdf_bullet(text: str, styles, dark=False) -> Paragraph:
    style = styles["bullet_dark" if dark else "bullet"]
    return Paragraph(f'<font color="#C79A3B">•</font>&nbsp;&nbsp;{_esc(text)}', style)


def _pdf_exp_block(exp: dict[str, Any], bullets: list[str], styles) -> list:
    header = Paragraph(_esc(exp["title"].upper()), styles["role"])
    company_line = Paragraph(
        f'<b><font color="#C79A3B">{_esc(exp["company"])}</font></b>'
        f'&nbsp;&nbsp;<font color="#5E6773">|&nbsp;&nbsp;{_esc(exp["location"])}&nbsp;&nbsp;|&nbsp;&nbsp;{_esc(exp["dates"])}</font>',
        styles["company"],
    )
    content = [header, company_line]
    content.extend(_pdf_bullet(item, styles) for item in bullets)
    return content


def _pdf_left_page1(styles) -> Table:
    items: list[Any] = []
    items.extend(_pdf_section("Executive Value", styles, dark=True))
    items.extend(_pdf_bullet(x, styles, dark=True) for x in PROFILE["executive_value"])
    items.extend(_pdf_section("Selected Partnerships", styles, dark=True))
    for label, value in PROFILE["selected_partnerships"]:
        items.append(Paragraph(_esc(label.upper()), styles["label_dark"]))
        items.append(Paragraph(_esc(value), styles["body_dark"]))
    items.extend(_pdf_section("Education", styles, dark=True))
    for label, value in PROFILE["education"]:
        items.append(Paragraph(_esc(label.upper()), styles["label_dark"]))
        items.append(Paragraph(_esc(value), styles["body_dark"]))
    items.extend(_pdf_section("Languages", styles, dark=True))
    items.extend(Paragraph(_esc(x), styles["body_dark"]) for x in PROFILE["languages"])
    table = Table([[items]], colWidths=[48 * mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#142B4A")),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("LINEAFTER", (0, 0), (-1, -1), 0.8, colors.HexColor("#C79A3B")),
    ]))
    return table


def _pdf_main_page1(result: dict[str, Any], styles) -> Table:
    items: list[Any] = []
    items.extend(_pdf_section("Executive Profile", styles))
    items.append(Paragraph(_esc(_profile_text(result)), styles["body"]))
    items.extend(_pdf_section("Signature Impact", styles))
    impact_rows = []
    for label, text in _signature_impact(result):
        impact_rows.append([Paragraph(f'<b><font color="#C79A3B">{_esc(label)}</font></b>&nbsp;&nbsp;{_esc(text)}', styles["impact"])])
    impact = Table(impact_rows, colWidths=[133 * mm])
    impact.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F3F5F7")),
        ("BOX", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9DEE5")),
        ("LINEABOVE", (0, 0), (-1, 0), 1.2, colors.HexColor("#C79A3B")),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5),
    ]))
    items.append(impact)
    items.extend(_pdf_section("Professional Experience", styles))
    current = dict(CURRENT_ROLE)
    items.extend(_pdf_exp_block(current, _current_bullets(result), styles))
    table = Table([[items]], colWidths=[136 * mm])
    table.setStyle(TableStyle([
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    return table


def _pdf_left_page2(result: dict[str, Any], styles) -> Table:
    items: list[Any] = []
    items.extend(_pdf_section("Core Capabilities", styles, dark=True))
    for label, skills in _capability_groups(result):
        items.append(Paragraph(_esc(label.upper()), styles["label_dark"]))
        items.append(Paragraph("<br/>".join(_esc(x) for x in skills), styles["body_dark"]))
    items.extend(_pdf_section("Tools", styles, dark=True))
    items.append(Paragraph("<br/>".join(_esc(x) for x in PROFILE["tools"]), styles["body_dark"]))
    items.extend(_pdf_section("Professional Development", styles, dark=True))
    selected = _list(result.get("selected_completed_courses"), limit=6)
    development = selected or PROFILE["professional_development"]
    items.extend(_pdf_bullet(x, styles, dark=True) for x in development)
    table = Table([[items]], colWidths=[48 * mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#142B4A")),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("LINEAFTER", (0, 0), (-1, -1), 0.8, colors.HexColor("#C79A3B")),
    ]))
    return table


def _pdf_main_page2(result: dict[str, Any], styles) -> Table:
    items: list[Any] = []
    items.extend(_pdf_section("Earlier Professional Experience", styles))
    for exp in OLDER_EXPERIENCE:
        items.extend(_pdf_exp_block(exp, exp["bullets"], styles))
    items.extend(_pdf_section("Leadership Positioning", styles))
    position = _clean(result.get("current_role_positioning_line")) or (
        "Best aligned to Brand Experience, Strategic Partnerships, Commercial Partnerships, "
        "Marketing Programs and Experiential Marketing leadership roles."
    )
    items.append(Paragraph(_esc(_shorten(position, 42)), styles["position"]))
    table = Table([[items]], colWidths=[136 * mm])
    table.setStyle(TableStyle([
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    return table


def _page_number(canvas, doc) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 6.5)
    canvas.setFillColor(colors.HexColor("#5E6773"))
    canvas.drawCentredString(A4[0] / 2, 5.5 * mm, f"ABDELRHMAN EL SHISHINY  •  EXECUTIVE RESUME  •  PAGE {doc.page}")
    canvas.restoreState()


def generate_premium_pdf(result: dict[str, Any], target_role: str = "") -> bytes:
    styles = _pdf_styles()
    headline = _headline(result, target_role)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=8 * mm,
        rightMargin=8 * mm,
        topMargin=7 * mm,
        bottomMargin=10 * mm,
        title=f'{PROFILE["name"]} Premium Executive Resume',
        author=PROFILE["name"],
    )
    story: list[Any] = []
    story.append(_pdf_header(styles, headline, compact=False))
    story.append(Spacer(1, 2.5))
    metric_cards = [_pdf_metric_card(value, label, styles) for value, label in _career_highlights(result)]
    metrics_table = Table([metric_cards], colWidths=[38.8 * mm] * 5, hAlign="CENTER")
    metrics_table.setStyle(TableStyle([
        ("LEFTPADDING", (0, 0), (-1, -1), 0.8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0.8),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(metrics_table)
    story.append(Spacer(1, 2.5))
    page1 = Table(
        [[_pdf_left_page1(styles), _pdf_main_page1(result, styles)]],
        colWidths=[49 * mm, 145 * mm],
        rowHeights=[222 * mm],
        hAlign="CENTER",
    )
    page1.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, 0), colors.HexColor("#142B4A")),
        ("LINEAFTER", (0, 0), (0, 0), 0.8, colors.HexColor("#C79A3B")),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(page1)
    story.append(PageBreak())
    story.append(_pdf_header(styles, headline, compact=True))
    story.append(Spacer(1, 2.5))
    page2 = Table(
        [[_pdf_left_page2(result, styles), _pdf_main_page2(result, styles)]],
        colWidths=[49 * mm, 145 * mm],
        rowHeights=[247 * mm],
        hAlign="CENTER",
    )
    page2.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, 0), colors.HexColor("#142B4A")),
        ("LINEAFTER", (0, 0), (0, 0), 0.8, colors.HexColor("#C79A3B")),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(page2)
    doc.build(story, onFirstPage=_page_number, onLaterPages=_page_number)
    return buffer.getvalue()
